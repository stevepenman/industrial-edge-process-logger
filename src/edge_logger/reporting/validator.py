"""Independent read-only XLSX validation before an email becomes eligible."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

if TYPE_CHECKING:
    from ..analysis.pipeline import ReportAnalysis

REQUIRED_SHEETS = (
    "Summary",
    "Temperature Summary",
    "Temperature Trends",
    "Process Trends",
    "Events - Exceptions",
    "Active Raw Data",
    "Constant Signals",
    "Source Manifest",
)


@dataclass(frozen=True, slots=True)
class WorkbookValidation:
    valid: bool
    errors: tuple[str, ...]


def validate_workbook(
    path: Path,
    expected_report_id: str | None = None,
    analysis: ReportAnalysis | None = None,
) -> WorkbookValidation:
    errors: list[str] = []
    if not path.exists() or path.stat().st_size == 0:
        return WorkbookValidation(False, ("Workbook is missing or empty",))
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            missing = set(REQUIRED_SHEETS).difference(workbook.sheetnames)
            if missing:
                errors.append(f"Missing mandatory worksheets: {', '.join(sorted(missing))}")
            if expected_report_id is not None and "Summary" in workbook.sheetnames:
                summary = workbook["Summary"]
                found = _key_values(summary).get("Report ID")
                if found != expected_report_id:
                    errors.append(f"Workbook report ID {found!r} does not match expected report ID")
            if expected_report_id is not None and "Source Manifest" in workbook.sheetnames:
                manifest_id = _key_values(workbook["Source Manifest"]).get("Report ID")
                if manifest_id != expected_report_id:
                    errors.append(
                        f"Source manifest report ID {manifest_id!r} does not match expected report ID"
                    )
            if analysis is not None:
                _validate_reconciliation(workbook, analysis, errors)
            _validate_sheet_structure(path, workbook.sheetnames, errors)
        finally:
            workbook.close()
    except (BadZipFile, InvalidFileException, KeyError, OSError, ValueError) as error:
        errors.append(f"Workbook cannot be reopened: {error}")
    return WorkbookValidation(not errors, tuple(errors))


def _key_values(sheet: Any) -> dict[str, object]:
    values: dict[str, object] = {}
    for row in sheet.iter_rows(min_col=1, max_col=8, values_only=True):
        for label_column, value_column in ((0, 1), (3, 4)):
            if isinstance(row[label_column], str):
                values[row[label_column]] = row[value_column]
    return values


def _validate_reconciliation(workbook: Any, analysis: ReportAnalysis, errors: list[str]) -> None:
    summary = _key_values(workbook["Summary"])
    manifest = _key_values(workbook["Source Manifest"])
    raw = workbook["Active Raw Data"]
    events = workbook["Events - Exceptions"]
    if summary.get("Raw record count") != analysis.record_count:
        errors.append("Summary raw record count does not reconcile")
    if manifest.get("Row count") != analysis.record_count:
        errors.append("Source manifest row count does not reconcile")
    if raw.max_row - _header_row(raw, "Source timestamp") != analysis.record_count:
        errors.append("Active Raw Data row count does not reconcile")
    expected_events = analysis.process_exception_count + analysis.data_quality_exception_count
    if events.max_row - _header_row(events, "Timestamp") != expected_events:
        errors.append("Events row count does not reconcile")
    if summary.get("Active signals") != len(analysis.active_profiles):
        errors.append("Summary active signal count does not reconcile")
    if summary.get("Constant signals") != len(analysis.constant_profiles):
        errors.append("Summary constant signal count does not reconcile")
    if summary.get("Process exceptions") != analysis.process_exception_count:
        errors.append("Summary process exception count does not reconcile")
    if summary.get("Data-quality exceptions") != analysis.data_quality_exception_count:
        errors.append("Summary data-quality exception count does not reconcile")
    if manifest.get("First timestamp") != (
        analysis.first_timestamp.isoformat() if analysis.first_timestamp else None
    ):
        errors.append("Source manifest first timestamp does not reconcile")
    if manifest.get("Last timestamp") != (
        analysis.last_timestamp.isoformat() if analysis.last_timestamp else None
    ):
        errors.append("Source manifest last timestamp does not reconcile")
    if manifest.get("First cycle") != analysis.first_cycle:
        errors.append("Source manifest first cycle does not reconcile")
    if manifest.get("Last cycle") != analysis.last_cycle:
        errors.append("Source manifest last cycle does not reconcile")
    for profile in analysis.active_profiles:
        if profile.is_clean_constant:
            errors.append(f"Classifier contradicts active signal {profile.signal_id}")
    if analysis.trends.get("temperature") and workbook["Temperature Trends"].max_row <= _header_row(
        workbook["Temperature Trends"], "Interval"
    ):
        errors.append("Temperature Trends unexpectedly has no trend rows")


def _header_row(sheet: Any, expected_first_header: str) -> int:
    """Return the one-based header row without depending on cosmetic title rows."""

    for row_index, row in enumerate(
        sheet.iter_rows(min_col=1, max_col=1, values_only=True), start=1
    ):
        if row[0] == expected_first_header:
            return row_index
    return 1


def _validate_sheet_structure(path: Path, names: list[str], errors: list[str]) -> None:
    """Inspect lightweight workbook XML rather than loading raw data into memory."""
    namespace = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
    filter_sheets = {
        "Temperature Summary",
        "Temperature Trends",
        "Process Trends",
        "Events - Exceptions",
        "Active Raw Data",
        "Constant Signals",
    }
    frozen_sheets = {
        "Temperature Summary",
        "Temperature Trends",
        "Process Trends",
        "Events - Exceptions",
        "Active Raw Data",
        "Constant Signals",
    }
    try:
        with ZipFile(path) as archive:
            for index, name in enumerate(names, start=1):
                has_filter, has_frozen_pane = _worksheet_features(archive, index)
                if name in filter_sheets and not has_filter:
                    errors.append(f"{name} is missing an autofilter")
                if name in frozen_sheets and not has_frozen_pane:
                    errors.append(f"{name} is missing a frozen pane")
            for chart_name in (
                entry for entry in archive.namelist() if entry.startswith("xl/charts/chart")
            ):
                chart = ElementTree.fromstring(archive.read(chart_name))
                formulas = [element.text or "" for element in chart.findall(".//c:f", namespace)]
                if not formulas:
                    errors.append(f"{chart_name} has no series ranges")
                if any("!" not in formula or "$" not in formula for formula in formulas):
                    errors.append(f"{chart_name} contains an invalid chart range")
    except (BadZipFile, ElementTree.ParseError, KeyError, OSError) as error:
        errors.append(f"Workbook structure cannot be inspected: {error}")


def _worksheet_features(archive: ZipFile, index: int) -> tuple[bool, bool]:
    """Inspect worksheet metadata without materialising a large raw-data XML worksheet."""
    has_filter = False
    has_frozen_pane = False
    with archive.open(f"xl/worksheets/sheet{index}.xml") as stream:
        for _, element in ElementTree.iterparse(stream, events=("end",)):
            tag = element.tag.rsplit("}", maxsplit=1)[-1]
            if tag == "autoFilter":
                has_filter = True
            elif tag == "pane" and element.attrib.get("state") == "frozen":
                has_frozen_pane = True
            if has_filter and has_frozen_pane:
                break
            element.clear()
    return has_filter, has_frozen_pane

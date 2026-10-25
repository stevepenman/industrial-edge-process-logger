"""Verify replay coverage, file hashes, receiver ingestion and encoded mail size offline."""

from __future__ import annotations

import argparse
import json
import sqlite3
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from edge_logger.email.sender import EmailConfiguration, send_report
from edge_logger.reporting.validator import validate_workbook
from edge_logger.storage.audit_db import AuditDatabase
from edge_logger.storage.raw_archive import archive_checksum
from edge_logger.storage.receiver import ingest_export


def verify(root: Path) -> dict[str, object]:
    database_path = root / "state" / "state.sqlite3"
    check = AuditDatabase.read_only_check(database_path)
    if not check["valid"]:
        raise ValueError(f"SQLite integrity check failed: {check}")
    database = sqlite3.connect(f"{database_path.resolve().as_uri()}?mode=ro", uri=True)
    database.row_factory = sqlite3.Row
    receiver = root / "receiver-verification.sqlite3"
    reports = []
    try:
        expected = int(database.execute("SELECT COUNT(*) FROM process_records").fetchone()[0])
        for source in database.execute("SELECT archive_path, sha256 FROM source_files"):
            if archive_checksum(Path(source["archive_path"])) != source["sha256"]:
                raise ValueError("Original source checksum mismatch")
        for row in database.execute(
            "SELECT * FROM report_periods WHERE workbook_path IS NOT NULL ORDER BY started_at"
        ):
            workbook = Path(row["workbook_path"])
            if archive_checksum(workbook) != row["workbook_sha256"]:
                raise ValueError("Workbook checksum mismatch")
            validation = validate_workbook(workbook, row["report_id"])
            if not validation.valid:
                raise ValueError(str(validation.errors))
            book = load_workbook(workbook, read_only=True)
            try:
                if book["Active Raw Data"].max_row - 1 != row["record_count"]:
                    raise ValueError("Workbook raw rows do not reconcile")
            finally:
                book.close()
            metadata = json.loads(
                database.execute(
                    "SELECT value FROM application_state WHERE key = ?",
                    (f"process_export:{row['id']}",),
                ).fetchone()[0]
            )
            export = Path(metadata["path"])
            if archive_checksum(export) != metadata["sha256"]:
                raise ValueError("Cycle export checksum mismatch")
            receipt = ingest_export(export, receiver)
            if receipt["record_count"] != row["record_count"]:
                raise ValueError("Receiver count does not match report count")
            if not ingest_export(export, receiver)["duplicate"]:
                raise ValueError("Receiver failed duplicate protection")
            encoded_sizes: list[int] = []
            with patch(
                "edge_logger.email.sender._send_message",
                side_effect=lambda _, message, sizes=encoded_sizes: sizes.append(
                    len(message.as_bytes())
                ),
            ):
                send_report(
                    EmailConfiguration(
                        "unused.invalid", 587, "logger@unused.invalid", ("qa@unused.invalid",)
                    ),
                    row["report_id"],
                    workbook,
                    "offline validation",
                    export,
                )
            if encoded_sizes[0] > 25 * 1024 * 1024:
                raise ValueError("Encoded email exceeds the default message size limit")
            reports.append(
                {
                    "report_id": row["report_id"],
                    "records": row["record_count"],
                    "encoded_email_bytes": encoded_sizes[0],
                    "receipt": receipt["receipt"],
                }
            )
        connection = sqlite3.connect(receiver)
        try:
            received = int(connection.execute("SELECT COUNT(*) FROM cycles").fetchone()[0])
        finally:
            connection.close()
        if received != expected or sum(item["records"] for item in reports) != expected:
            raise ValueError("End-to-end cycle coverage does not reconcile")
        result = {
            "valid": True,
            "canonical_records": expected,
            "receiver_records": received,
            "reports": reports,
            "sqlite_integrity": check["valid"],
            "network_used": False,
        }
        (root / "verification.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
        return result
    finally:
        database.close()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("replay_directory", type=Path)
    print(json.dumps(verify(parser.parse_args().replay_directory), indent=2))


if __name__ == "__main__":
    main()

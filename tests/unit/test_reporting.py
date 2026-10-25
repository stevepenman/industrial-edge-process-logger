"""Synthetic fixtures only; no operational datasets are required."""

from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

from edge_logger.analysis.pipeline import analyse_records, iter_events
from edge_logger.models import ProcessRecord, QualityStatus, SignalValue
from edge_logger.reporting.charts import series_colour
from edge_logger.reporting.generator import _recompress_workbook, _temperature_axis, generate_report
from edge_logger.reporting.validator import validate_workbook


def test_transactional_report_contains_mandatory_sheets(tmp_path: Path) -> None:
    timestamp = datetime(2026, 8, 20, 12, tzinfo=UTC)
    records = [
        ProcessRecord(
            "DEMO-01",
            timestamp,
            timestamp,
            "1",
            "fixture:2",
            (
                SignalValue("tcd05", "TCD05", 0, "C", "float", QualityStatus.GOOD, timestamp),
                SignalValue(
                    "cycle_time", "Cycle time", 9.5, "s", "float", QualityStatus.GOOD, timestamp
                ),
            ),
        ),
        ProcessRecord(
            "DEMO-01",
            timestamp,
            timestamp,
            "2",
            "fixture:3",
            (
                SignalValue("tcd05", "TCD05", 0, "C", "float", QualityStatus.GOOD, timestamp),
                SignalValue(
                    "cycle_time", "Cycle time", 10.5, "s", "float", QualityStatus.GOOD, timestamp
                ),
            ),
        ),
    ]
    report_id = "f5e93f44-15e0-4b33-8ce3-b74a0edbead0"
    path, checksum = generate_report(
        records, tmp_path, report_id, "DEMO-01", "Europe/London", {"configuration_version": 1}
    )
    assert path.exists()
    assert checksum
    validation = validate_workbook(path, report_id)
    assert validation.valid, validation.errors


def test_active_raw_data_excludes_clean_constant_signals(tmp_path: Path) -> None:
    timestamp = datetime(2026, 8, 20, 12, tzinfo=UTC)
    records = [
        ProcessRecord(
            "DEMO-01",
            timestamp + timedelta(minutes=index),
            timestamp + timedelta(minutes=index),
            str(index + 1),
            f"fixture:{index + 1}",
            (
                SignalValue(
                    "constant",
                    "Constant signal",
                    7.5,
                    "bar",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "variable",
                    "Variable signal",
                    10.0 + index,
                    "bar",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
            ),
        )
        for index in range(2)
    ]
    report, _ = generate_report(
        records,
        tmp_path,
        "f5e93f44-15e0-4b33-8ce3-b74a0edbead0",
        "DEMO-01",
        "Europe/London",
        {},
    )
    workbook = load_workbook(report, read_only=True, data_only=False)
    try:
        raw = workbook["Active Raw Data"]
        assert [raw.cell(1, column).value for column in range(1, 7)] == [
            "Source timestamp",
            "Receive timestamp",
            "Date",
            "Time",
            "Cycle ID",
            "Variable signal (bar)",
        ]
        assert raw.max_column == 6
        assert [raw.cell(row, 6).value for row in (2, 3)] == [10.0, 11.0]
    finally:
        workbook.close()


def test_workbook_recompression_preserves_contents_and_reduces_repetitive_data(
    tmp_path: Path,
) -> None:
    workbook = tmp_path / "fixture.xlsx"
    payload = b"<worksheet><row><cell>Repeated process data</cell></row></worksheet>" * 10_000
    with ZipFile(workbook, "w", compression=ZIP_DEFLATED, compresslevel=1) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", payload)
    original_size = workbook.stat().st_size

    _recompress_workbook(workbook)

    assert workbook.stat().st_size < original_size
    with ZipFile(workbook) as archive:
        assert archive.read("xl/worksheets/sheet1.xml") == payload


def test_temperature_chart_tabs_use_cycle_axis(tmp_path: Path) -> None:
    first = datetime(2026, 8, 20, 12, tzinfo=UTC)
    second = datetime(2026, 8, 21, 12, tzinfo=UTC)
    records = [
        ProcessRecord(
            "DEMO-01",
            first,
            first,
            "1",
            "fixture:1",
            (
                SignalValue("feed", "Feed", 50.0, None, "float", QualityStatus.GOOD, first),
                SignalValue("zone", "Zone", 200.0, None, "float", QualityStatus.GOOD, first),
                SignalValue("water", "Water", 20.0, None, "float", QualityStatus.GOOD, first),
                SignalValue("oil", "Oil", 30.0, None, "float", QualityStatus.GOOD, first),
            ),
        ),
        ProcessRecord(
            "DEMO-01",
            second,
            second,
            "2",
            "fixture:2",
            (
                SignalValue("feed", "Feed", 51.0, None, "float", QualityStatus.GOOD, second),
                SignalValue("zone", "Zone", 201.0, None, "float", QualityStatus.GOOD, second),
                SignalValue("water", "Water", 21.0, None, "float", QualityStatus.GOOD, second),
                SignalValue("oil", "Oil", 31.0, None, "float", QualityStatus.GOOD, second),
            ),
        ),
    ]
    report, _ = generate_report(
        records,
        tmp_path,
        "f5e93f44-15e0-4b33-8ce3-b74a0edbead0",
        "DEMO-01",
        "Europe/London",
        {},
        {
            "feed": {
                "category": "temperature",
                "trend": True,
                "chart_group": "barrel",
                "chart_axis": "secondary",
            },
            "zone": {"category": "temperature", "trend": True, "chart_group": "barrel"},
            "water": {"category": "temperature", "trend": True, "chart_group": "tcd"},
            "oil": {"category": "temperature", "trend": True, "chart_group": "oil"},
        },
    )
    workbook = load_workbook(report, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Summary",
            "Temperature Summary",
            "Temperature Trends",
            "Barrel Temperatures",
            "Water (TCD) Temperatures",
            "Oil Temperature",
            "Process Trends",
            "Events - Exceptions",
            "Active Raw Data",
            "Constant Signals",
            "Source Manifest",
        ]
        assert workbook["Temperature Trends"]["A1"].value == "TEMPERATURE TRENDS"
        assert workbook["Process Trends"]["A1"].value == "PROCESS TRENDS"
        assert workbook["Barrel Temperatures"]["A1"].value == "BARREL TEMPERATURES"
        temperature_headers = next(
            workbook["Temperature Summary"].iter_rows(min_row=4, max_row=4, values_only=True)
        )
        assert temperature_headers == (
            "Signal",
            "Unit",
            "Records",
            "Valid readings",
            "Missing/NULL",
            "Unique values",
            "Minimum",
            "Average",
            "Std. dev.",
            "Maximum",
            "Range/span",
            "Zero count",
            "Low limit",
            "High limit",
            "Low violations",
            "High violations",
            "Status",
        )
    finally:
        workbook.close()
    with ZipFile(report) as workbook:
        chart_xml = workbook.read("xl/charts/chart1.xml")
        assert len([name for name in workbook.namelist() if name.startswith("xl/charts/")]) == 3
    assert b"Cycle number" in chart_xml
    assert b"$C$5:$C$6" in chart_xml
    assert b"Temperature change" not in chart_xml
    assert b"majorTimeUnit" not in chart_xml


def test_process_chart_tabs_do_not_mix_engineering_units(tmp_path: Path) -> None:
    timestamp = datetime(2026, 8, 20, 12, tzinfo=UTC)
    records = [
        ProcessRecord(
            "DEMO-01",
            timestamp + timedelta(minutes=index),
            timestamp + timedelta(minutes=index),
            str(index + 1),
            f"fixture:{index + 1}",
            (
                SignalValue(
                    "pressure",
                    "Pressure",
                    100.0 + index,
                    "bar",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "integral",
                    "Integral",
                    500.0 + index,
                    "kPa",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "position",
                    "Position",
                    10.0 + index,
                    "mm",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "duration", "Duration", 2.0 + index, "s", "float", QualityStatus.GOOD, timestamp
                ),
            ),
        )
        for index in range(2)
    ]
    registry = {
        "pressure": {
            "category": "process",
            "trend": True,
            "chart_group": "injection_pressure",
            "unit": "bar",
        },
        "integral": {
            "category": "process",
            "trend": True,
            "chart_group": "injection_pressure",
            "unit": "kPa",
        },
        "position": {"category": "process", "trend": True, "chart_group": "screw", "unit": "mm"},
        "duration": {"category": "process", "trend": True, "chart_group": "screw", "unit": "s"},
    }
    report, _ = generate_report(
        records,
        tmp_path,
        "f5e93f44-15e0-4b33-8ce3-b74a0edbead0",
        "DEMO-01",
        "Europe/London",
        {},
        registry,
    )
    workbook = load_workbook(report, read_only=True, data_only=False)
    try:
        assert {
            "Injection pressure (bar)",
            "Injection pressure (kPa)",
            "Screw-back process (mm)",
            "Screw-back process (s)",
        }.issubset(workbook.sheetnames)
    finally:
        workbook.close()


def test_summary_includes_energy_performance(tmp_path: Path) -> None:
    timestamp = datetime(2026, 8, 20, 12, tzinfo=UTC)
    records = [
        ProcessRecord(
            "DEMO-01",
            timestamp + timedelta(minutes=index),
            timestamp + timedelta(minutes=index),
            str(index + 1),
            f"fixture:{index + 1}",
            (
                SignalValue(
                    "energy_cycle_drives_wh",
                    "Drives energy per cycle",
                    140.0 + index * 10,
                    "Wh",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "energy_cycle_heating_wh",
                    "Heating energy per cycle",
                    10.0 + index * 2,
                    "Wh",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "energy_cycle_total_wh",
                    "Total energy per cycle",
                    150.0 + index * 12,
                    "Wh",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "energy_per_kg_drives",
                    "Drives specific energy",
                    0.60 + index * 0.01,
                    "kWh/kg",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "energy_per_kg_heating",
                    "Heating specific energy",
                    0.04 + index * 0.01,
                    "kWh/kg",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "energy_per_kg_total",
                    "Total specific energy",
                    0.64 + index * 0.02,
                    "kWh/kg",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
            ),
        )
        for index in range(2)
    ]
    report, _ = generate_report(
        records, tmp_path, "f5e93f44-15e0-4b33-8ce3-b74a0edbead0", "DEMO-01", "Europe/London", {}
    )
    workbook = load_workbook(report, read_only=True, data_only=False)
    try:
        summary = workbook["Summary"]
        assert summary["D16"].value == "Energy Performance"
        assert summary["E17"].value == 145.0
        assert summary["E23"].value == 0
    finally:
        workbook.close()


def test_chart_y_axis_matches_the_reported_extrema() -> None:
    assert _temperature_axis([199.82, 199.91, 200.07]) == {
        "min": 199.82,
        "max": 200.07,
        "major_unit": 0.05,
    }


def test_configured_tcd_channels_have_distinct_chart_colours() -> None:
    colours = {series_colour(f"tcd_{number:02d}_temperature") for number in range(1, 13)}
    assert len(colours) == 12


def test_report_table_headers_use_controlled_readable_widths(tmp_path: Path) -> None:
    timestamp = datetime(2026, 8, 20, 12, tzinfo=UTC)
    records = [
        ProcessRecord(
            "DEMO-01",
            timestamp,
            timestamp,
            "1",
            "fixture:1",
            (
                SignalValue(
                    "temperature",
                    "Long configured temperature name",
                    200.0,
                    "C",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "pressure",
                    "Maximum injection pressure",
                    1234.0,
                    "bar",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
            ),
        ),
        ProcessRecord(
            "DEMO-01",
            timestamp + timedelta(minutes=5),
            timestamp + timedelta(minutes=5),
            "2",
            "fixture:2",
            (
                SignalValue(
                    "temperature",
                    "Long configured temperature name",
                    201.0,
                    "C",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "pressure",
                    "Maximum injection pressure",
                    1235.0,
                    "bar",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
            ),
        ),
    ]
    report, _ = generate_report(
        records,
        tmp_path,
        "f5e93f44-15e0-4b33-8ce3-b74a0edbead0",
        "DEMO-01",
        "Europe/London",
        {},
        {"temperature": {"category": "temperature"}, "pressure": {"category": "process"}},
    )
    workbook = load_workbook(report, read_only=False, data_only=False)
    try:
        assert workbook["Summary"].column_dimensions["C"].width >= 13
        assert workbook["Summary"].column_dimensions["F"].width >= 22
        assert workbook["Temperature Summary"].column_dimensions["D"].width >= 16
        assert workbook["Events - Exceptions"].column_dimensions["G"].width >= 22
        assert 28 <= workbook["Active Raw Data"].column_dimensions["F"].width <= 29
        assert workbook["Active Raw Data"].row_dimensions[1].height == 45
    finally:
        workbook.close()


def test_report_opening_print_and_raw_units_metadata_presentation(tmp_path: Path) -> None:
    timestamp = datetime(2026, 8, 20, 12, tzinfo=UTC)
    records = [
        ProcessRecord(
            "DEMO-01",
            None,
            timestamp,
            None,
            "fixture:units",
            (
                SignalValue(
                    "temperature",
                    "Barrel temperature",
                    "C",
                    "C",
                    "string",
                    QualityStatus.GOOD,
                    None,
                ),
                SignalValue(
                    "pressure",
                    "Injection pressure",
                    "bar",
                    "bar",
                    "string",
                    QualityStatus.GOOD,
                    None,
                ),
            ),
        ),
        ProcessRecord(
            "DEMO-01",
            timestamp,
            timestamp,
            "1",
            "fixture:1",
            (
                SignalValue(
                    "temperature",
                    "Barrel temperature",
                    200.0,
                    "C",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "pressure",
                    "Injection pressure",
                    1234.0,
                    "bar",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
            ),
        ),
        ProcessRecord(
            "DEMO-01",
            timestamp + timedelta(minutes=5),
            timestamp + timedelta(minutes=5),
            "2",
            "fixture:2",
            (
                SignalValue(
                    "temperature",
                    "Barrel temperature",
                    201.0,
                    "C",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
                SignalValue(
                    "pressure",
                    "Injection pressure",
                    1235.0,
                    "bar",
                    "float",
                    QualityStatus.GOOD,
                    timestamp,
                ),
            ),
        ),
    ]
    report, _ = generate_report(
        records,
        tmp_path,
        "f5e93f44-15e0-4b33-8ce3-b74a0edbead0",
        "DEMO-01",
        "Europe/London",
        {},
        {"temperature": {"category": "temperature"}, "pressure": {"category": "process"}},
    )
    workbook = load_workbook(report, read_only=False, data_only=False)
    try:
        summary = workbook["Summary"]
        temperature_summary = workbook["Temperature Summary"]
        trends = workbook["Temperature Trends"]
        events = workbook["Events - Exceptions"]
        raw = workbook["Active Raw Data"]
        manifest = workbook["Source Manifest"]
        assert workbook.active.title == "Summary"
        assert summary.freeze_panes is None
        assert manifest.freeze_panes is None
        assert summary.print_area == f"'Summary'!$A$1:$H${summary.max_row}"
        assert summary.sheet_view.zoomScale in {None, 100}
        assert temperature_summary.sheet_view.zoomScale == 90
        assert events.sheet_view.zoomScale == 80
        assert raw.sheet_view.zoomScale == 75
        assert raw.freeze_panes == "F2"
        assert temperature_summary.print_title_rows == "$4:$4"
        assert trends.print_title_rows == "$4:$4"
        assert events.print_title_rows == "$4:$4"
        assert raw.print_title_rows == "$1:$1"
        assert summary.page_setup.orientation == "portrait"
        assert events.page_setup.orientation == "landscape"
        assert summary.sheet_properties.pageSetUpPr.fitToPage
        assert events.sheet_properties.pageSetUpPr.fitToPage
        assert summary.page_setup.fitToWidth in {None, 1}
        assert events.page_setup.fitToWidth in {None, 1}
        assert events.page_setup.fitToHeight == 0
        assert summary.oddFooter.left.text == "DEMO-01 Process Report"
        assert summary.oddFooter.center.text == "Report ID: f5e93f44"
        assert summary.oddFooter.right.text == "Page &P of &N"
        assert raw["F2"].value == "C"
        assert raw["B2"].value is None
        assert raw["F2"].font.italic
        assert raw["F2"].fill.fgColor.rgb.endswith("EAF2F8")
        assert raw.row_dimensions[2].height == 14
        assert raw["A1"].fill.fgColor.rgb.endswith("D9EAF7")
        assert raw["F1"].fill.fgColor.rgb.endswith("17365D")
    finally:
        workbook.close()


def test_workbook_values_are_preserved_independently_of_presentation(tmp_path: Path) -> None:
    """Report styling may move headers, but it must not change report semantics."""

    first = datetime(2026, 8, 20, 12, tzinfo=UTC)
    registry: dict[str, dict[str, object]] = {
        "feed": {
            "category": "temperature",
            "trend": True,
            "chart_group": "barrel",
            "chart_axis": "secondary",
        },
        "zone": {
            "category": "temperature",
            "trend": True,
            "chart_group": "barrel",
            "low_limit": 199.0,
        },
        "water": {"category": "temperature", "trend": True, "chart_group": "tcd"},
        "oil": {"category": "temperature", "trend": True, "chart_group": "oil"},
        "cycle_time": {"category": "process", "exception_rule": "long_cycle"},
        "bad_parts": {"category": "quality", "exception_rule": "bad_part"},
    }
    values = (
        (50.0, 200.0, 20.0, 30.0, 10.0, 0.0),
        (51.0, 201.0, 21.0, 31.0, 13.0, 1.0),
        (52.0, 202.0, 22.0, 32.0, 11.0, 0.0),
    )
    records = [
        ProcessRecord(
            "DEMO-01",
            first + timedelta(minutes=index * 5),
            first + timedelta(minutes=index * 5),
            cycle_id,
            f"fixture:{cycle_id}",
            (
                SignalValue("feed", "Feed", row[0], "C", "float", QualityStatus.GOOD, first),
                SignalValue("zone", "Zone", row[1], "C", "float", QualityStatus.GOOD, first),
                SignalValue("water", "Water", row[2], "C", "float", QualityStatus.GOOD, first),
                SignalValue("oil", "Oil", row[3], "C", "float", QualityStatus.GOOD, first),
                SignalValue(
                    "cycle_time", "Cycle time", row[4], "s", "float", QualityStatus.GOOD, first
                ),
                SignalValue(
                    "bad_parts", "Bad parts", row[5], None, "float", QualityStatus.GOOD, first
                ),
            ),
        )
        for index, (cycle_id, row) in enumerate(zip(("100", "101", "103"), values, strict=True))
    ]
    report_id = "f5e93f44-15e0-4b33-8ce3-b74a0edbead0"
    source_manifest: dict[str, object] = {
        "acquisition_hours": 120,
        "application_version": "test-version",
        "configuration_version": 7,
        "original_filenames": "deterministic.csv",
    }
    rules: dict[str, object] = {"long_cycle_seconds": 12.0}
    analysis = analyse_records(
        lambda: iter(records), registry, "Europe/London", 5, rules, trend_bucket_count=30
    )
    report, _ = generate_report(
        records,
        tmp_path,
        report_id,
        "DEMO-01",
        "Europe/London",
        source_manifest,
        registry,
        rules,
    )

    workbook = load_workbook(report, read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == [
            "Summary",
            "Temperature Summary",
            "Temperature Trends",
            "Barrel Temperatures",
            "Water (TCD) Temperatures",
            "Oil Temperature",
            "Process Trends",
            "Events - Exceptions",
            "Active Raw Data",
            "Constant Signals",
            "Source Manifest",
        ]
        summary = _key_values(workbook["Summary"])
        assert summary["Machine"] == "DEMO-01"
        assert summary["Report ID"] == report_id
        assert summary["Raw record count"] == analysis.record_count
        assert summary["First cycle"] == analysis.first_cycle
        assert summary["Last cycle"] == analysis.last_cycle
        assert summary["Unique cycles"] == analysis.unique_cycles
        assert summary["Missing cycles"] == analysis.missing_cycles
        assert summary["Cycles flagged bad"] == analysis.bad_part_count
        assert summary["Long cycles"] == analysis.long_cycle_count
        assert summary["Process exceptions"] == analysis.process_exception_count
        assert summary["Data-quality exceptions"] == analysis.data_quality_exception_count

        temperature_sheet = workbook["Temperature Summary"]
        assert _row_values(temperature_sheet, 4, 17) == (
            "Signal",
            "Unit",
            "Records",
            "Valid readings",
            "Missing/NULL",
            "Unique values",
            "Minimum",
            "Average",
            "Std. dev.",
            "Maximum",
            "Range/span",
            "Zero count",
            "Low limit",
            "High limit",
            "Low violations",
            "High violations",
            "Status",
        )
        expected_profiles = [
            profile for profile in analysis.profiles.values() if profile.category == "temperature"
        ]
        actual_profiles = [
            _row_values(temperature_sheet, row, 17) for row in range(5, 5 + len(expected_profiles))
        ]
        assert actual_profiles == [
            (
                profile.display_name,
                "°C" if profile.unit == "C" else profile.unit,
                profile.rows,
                profile.valid_readings,
                profile.missing_readings,
                profile.unique_values,
                profile.minimum,
                profile.average,
                profile.standard_deviation,
                profile.maximum,
                profile.maximum - profile.minimum
                if profile.maximum is not None and profile.minimum is not None
                else None,
                profile.zero_count,
                profile.metadata.get("low_limit"),
                profile.metadata.get("high_limit"),
                profile.low_limit_violations,
                profile.high_limit_violations,
                profile.status,
            )
            for profile in expected_profiles
        ]

        trends_sheet = workbook["Temperature Trends"]
        trend_columns = 5 + len(expected_profiles)
        actual_trends = [
            _row_values(trends_sheet, row, trend_columns)
            for row in range(5, 5 + len(analysis.trends["temperature"]))
        ]
        assert actual_trends == [
            (
                trend.interval,
                trend.cycle_start,
                trend.cycle_end,
                trend.start_timestamp.astimezone(ZoneInfo("Europe/London")).replace(tzinfo=None)
                if trend.start_timestamp
                else None,
                trend.end_timestamp.astimezone(ZoneInfo("Europe/London")).replace(tzinfo=None)
                if trend.end_timestamp
                else None,
                *(trend.values.get(profile.signal_id) for profile in expected_profiles),
            )
            for trend in analysis.trends["temperature"]
        ]

        events_sheet = workbook["Events - Exceptions"]
        actual_events = [
            _row_values(events_sheet, row, 9) for row in range(5, events_sheet.max_row + 1)
        ]
        expected_events = list(iter_events(lambda: iter(records), analysis, rules))
        london = ZoneInfo("Europe/London")
        assert actual_events == [
            (
                event.timestamp.astimezone(london).replace(tzinfo=None),
                event.cycle_id,
                event.category,
                event.event_type,
                event.signal_id,
                event.value,
                event.threshold,
                event.severity,
                event.detail,
            )
            for event in expected_events
        ]

        raw_sheet = workbook["Active Raw Data"]
        active_profiles = list(analysis.active_profiles)
        assert _row_values(raw_sheet, 1, 5 + len(active_profiles)) == (
            "Source timestamp",
            "Receive timestamp",
            "Date",
            "Time",
            "Cycle ID",
            *(
                f"{profile.display_name} ({'°C' if profile.unit == 'C' else profile.unit})"
                if profile.unit
                else profile.display_name
                for profile in active_profiles
            ),
        )
        for row_index, record in enumerate(records, start=2):
            local = (
                record.source_timestamp.astimezone(london).replace(tzinfo=None)
                if record.source_timestamp
                else None
            )
            source_values = {signal.signal_id: signal.value for signal in record.signals}
            assert _row_values(raw_sheet, row_index, 5 + len(active_profiles)) == (
                local,
                local,
                local.date().isoformat() if local else None,
                local.time().isoformat(timespec="seconds") if local else None,
                record.cycle_id,
                *(source_values.get(profile.signal_id) for profile in active_profiles),
            )

        manifest = _key_values(workbook["Source Manifest"])
        assert manifest["Report ID"] == report_id
        assert manifest["Row count"] == analysis.record_count
        assert manifest["First cycle"] == analysis.first_cycle
        assert manifest["Last cycle"] == analysis.last_cycle
        assert manifest["Original filename"] == "deterministic.csv"
        assert not any(
            cell.data_type == "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
    finally:
        workbook.close()


def _key_values(sheet: Any) -> dict[str, object]:
    values: dict[str, object] = {}
    for row in sheet.iter_rows(min_col=1, max_col=8, values_only=True):
        for label_column, value_column in ((0, 1), (3, 4)):
            if isinstance(row[label_column], str):
                values[row[label_column]] = row[value_column]
    return values


def _row_values(sheet: Any, row: int, columns: int) -> tuple[object, ...]:
    return tuple(next(sheet.iter_rows(min_row=row, max_row=row, max_col=columns, values_only=True)))

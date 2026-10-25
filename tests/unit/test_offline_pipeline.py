"""Synthetic fixtures only; no operational datasets are required."""

import asyncio
from pathlib import Path

from openpyxl import load_workbook

from edge_logger.config import load_config
from edge_logger.runtime import LoggerRuntime
from edge_logger.storage.audit_db import AuditDatabase


def test_import_retains_source_and_builds_validated_report(tmp_path: Path) -> None:
    source = tmp_path / "demo.tsv"
    source.write_text(
        "Date\tTime\tCycles\tCycle time\tBad parts\tFeed zone actual temp.\tOiltemperature\n"
        "12.08.26\t18:15:00\t1\t9,0\t0\t50,0\t30,0\n"
        "12.08.26\t18:20:00\t2\t15,0\t1\t51,0\t31,0\n"
        "12.08.26\t18:25:00\t4\t9,0\t0\t50,0\t30,0\n",
        encoding="utf-16",
    )
    config_file = tmp_path / "config.yaml"
    lines = [
        "machine:",
        "  name: DEMO-01",
        "  timezone: Europe/London",
        "application:",
        "  config_version: 1",
        "reporting:",
        "  trigger_mode: acquisition_hours",
        "  acquisition_hours: 120",
        "  maximum_calendar_days: null",
        "  trend_bucket_minutes: 5",
        "  generate_if_due_on_startup: true",
        "  minimum_records: 1",
        "retention:",
        "  days_after_successful_email: 30",
        "ups:",
        "  shutdown_delay_seconds: 90",
        "rules:",
        "  long_cycle_seconds: 12.0",
        "opcua:",
        "  endpoint: null",
        "email:",
        "  enabled: false",
        "paths:",
        f"  raw: {tmp_path / 'raw'}",
        f"  reports: {tmp_path / 'reports'}",
        f"  state: {tmp_path / 'state'}",
        "signals:",
        "  cycle_time:",
        "    source_name: Cycle time",
        "    category: process",
        "    unit: s",
        "    expected_datatype: float",
        "    trend: true",
        "    exception_rule: long_cycle",
        "  bad_parts:",
        "    source_name: Bad parts",
        "    category: quality",
        "    expected_datatype: float",
        "    exception_rule: bad_part",
        "  feed_temperature:",
        "    source_name: Feed zone actual temp.",
        "    category: temperature",
        "    unit: C",
        "    expected_datatype: float",
        "    trend: true",
        "    chart_group: barrel",
        "  oil_temperature:",
        "    source_name: Oiltemperature",
        "    category: temperature",
        "    unit: C",
        "    expected_datatype: float",
        "    trend: true",
        "    chart_group: oil",
    ]
    config_file.write_text("\n".join(lines), encoding="utf-8")
    config = load_config(config_file)
    database = AuditDatabase(config.state_path / "state.sqlite3")
    try:
        runtime = LoggerRuntime(config, database)
        result = asyncio.run(runtime.import_file(source, force_report=True))
        assert result.imported_records == 3
        assert not result.duplicate
        assert len(result.sealed_report_ids) == 1
        period = database.report_period(result.sealed_report_ids[0])
        assert period["state"] == "REPORT_VALIDATED"
        assert Path(period["workbook_path"]).is_file()
        assert len(database.source_files_for_period(period["id"])) == 1
        assert len(database.raw_segments_for_period(period["id"])) == 1
        workbook = load_workbook(period["workbook_path"], read_only=True, data_only=True)
        try:
            assert workbook.sheetnames == [
                "Summary",
                "Temperature Summary",
                "Temperature Trends",
                "Barrel Temperatures",
                "Oil Temperature",
                "Process Trends",
                "Events - Exceptions",
                "Active Raw Data",
                "Constant Signals",
                "Source Manifest",
            ]
            assert workbook["Active Raw Data"].max_row == 4
            assert workbook["Source Manifest"]["B6"].value == result.sealed_report_ids[0]
            assert workbook["Summary"]["E6"].value == "N/A – historical import"
            manifest_values = {
                row[0]: row[1]
                for row in workbook["Source Manifest"].iter_rows(
                    min_col=1, max_col=2, values_only=True
                )
                if isinstance(row[0], str)
            }
            assert manifest_values["Acquisition provider"] == "File import"
        finally:
            workbook.close()
        duplicate = asyncio.run(runtime.import_file(source))
        assert duplicate.duplicate
        assert duplicate.imported_records == 0
    finally:
        database.close()

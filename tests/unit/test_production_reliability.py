"""Synthetic fixtures only; no operational datasets are required."""

import asyncio
import gzip
import json
import sqlite3
import subprocess
import sys
import threading
import time
from dataclasses import replace
from datetime import UTC, datetime, timedelta
from email.message import EmailMessage
from pathlib import Path
from types import SimpleNamespace

import pytest
from asyncua import ua

import edge_logger.runtime as runtime_module
from edge_logger.acquisition.base import AcquisitionProvider
from edge_logger.acquisition.file_export import FileExportImporter
from edge_logger.acquisition.opcua import (
    OpcUaConnectionSettings,
    OpcUaNode,
    OpcUaProvider,
    _quality,
)
from edge_logger.analysis.pipeline import analyse_records
from edge_logger.config import ConfigurationError, load_config
from edge_logger.email.outbox import EmailOutboxProcessor
from edge_logger.email.sender import EmailConfiguration, PartialDeliveryError, _send_message
from edge_logger.models import (
    EmailState,
    ProcessRecord,
    QualityStatus,
    ReportPeriodState,
    SignalValue,
)
from edge_logger.runtime import LoggerRuntime
from edge_logger.storage.audit_db import AuditDatabase
from edge_logger.storage.lock import AcquisitionLock
from edge_logger.storage.raw_archive import RawArchiveWriter, archive_checksum, read_records
from edge_logger.storage.receiver import ingest_export
from edge_logger.ups.command import CommandUpsProvider, request_host_shutdown


def _config(root: Path):
    base = load_config(Path("config/config.example.yaml"))
    return replace(
        base,
        raw_path=root / "raw",
        reports_path=root / "reports",
        state_path=root / "state",
        raw={
            **base.raw,
            "operations": {"automatic_retention": False, "minimum_free_disk_bytes": 1},
        },
    )


def _record(cycle: int = 1) -> ProcessRecord:
    timestamp = datetime(2026, 10, 24, 12, tzinfo=UTC) + timedelta(seconds=cycle)
    return ProcessRecord(
        "DEMO-01",
        timestamp,
        datetime.now(UTC),
        str(cycle),
        "fixture",
        (
            SignalValue(
                "pressure", "Pressure", 100 + cycle, "bar", "float", QualityStatus.GOOD, timestamp
            ),
        ),
    )


def _notification(timestamp: datetime, status=None):
    return SimpleNamespace(
        monitored_item=SimpleNamespace(
            Value=SimpleNamespace(
                SourceTimestamp=timestamp,
                ServerTimestamp=timestamp,
                StatusCode=status or ua.StatusCode(),
            )
        )
    )


def _node(identifier: str):
    return SimpleNamespace(nodeid=SimpleNamespace(to_string=lambda: identifier))


def _provider(**settings):
    return OpcUaProvider(
        OpcUaConnectionSettings("opc.tcp://127.0.0.1:4840", **settings),
        (
            OpcUaNode("pressure", "ns=2;i=1", "Pressure", "bar", "float"),
            OpcUaNode("cycle", "ns=2;i=2", "Cycle", None, "int", is_cycle_id=True),
        ),
    )


@pytest.mark.parametrize(
    ("code", "expected"),
    [
        (ua.StatusCodes.Good, QualityStatus.GOOD),
        (ua.StatusCodes.BadNoCommunication, QualityStatus.BAD),
        (ua.StatusCodes.UncertainLastUsableValue, QualityStatus.UNCERTAIN),
    ],
)
def test_real_numeric_opcua_status(code, expected):
    assert _quality(ua.StatusCode(code)) is expected
    assert _quality(None) is QualityStatus.MISSING


def test_cycle_before_process_update_waits_for_matching_timestamp_and_deduplicates():
    provider = _provider()
    earlier = _record(1).source_timestamp
    later = _record(2).source_timestamp

    async def exercise():
        await provider._data_change(_node("ns=2;i=1"), 100.0, _notification(earlier))
        await provider._data_change(_node("ns=2;i=2"), 1, _notification(earlier))
        provider._records.get_nowait()
        await provider._data_change(_node("ns=2;i=2"), 2, _notification(later))
        assert provider._records.empty()
        await provider._data_change(
            _node("ns=2;i=2"), 2, _notification(later + timedelta(seconds=1))
        )
        assert provider._records.empty()
        await provider._data_change(_node("ns=2;i=1"), 200.0, _notification(later))
        result = provider._records.get_nowait()
        assert next(value.value for value in result.signals if value.signal_id == "pressure") == 200
        assert result.production_context["snapshot_complete"]
        await provider._data_change(_node("ns=2;i=2"), 2, _notification(later))
        assert provider._records.empty()

    asyncio.run(exercise())


def test_expired_snapshot_marks_missing_instead_of_reusing_old_pressure():
    provider = _provider()

    async def exercise():
        await provider._data_change(
            _node("ns=2;i=1"), 100.0, _notification(_record(1).source_timestamp)
        )
        await provider._data_change(
            _node("ns=2;i=2"), 2, _notification(_record(2).source_timestamp)
        )
        provider._flush_pending(force=True)
        result = provider._records.get_nowait()
        pressure = next(value for value in result.signals if value.signal_id == "pressure")
        assert pressure.value is None and pressure.quality is QualityStatus.MISSING
        assert result.production_context["snapshot_complete"] is False

    asyncio.run(exercise())


def test_subscription_requests_timestamp_notifications_and_server_queue():
    provider = _provider(monitored_item_queue_size=123)
    captured = []

    class Node:
        def __init__(self, node_id):
            self.nodeid = ua.NodeId.from_string(node_id)

        async def read_data_value(self):
            return ua.DataValue()

    class Subscription:
        async def create_monitored_items(self, requests):
            captured.extend(requests)
            return [1, 2]

    asyncio.run(provider._subscribe(SimpleNamespace(get_node=Node), Subscription()))
    assert all(request.RequestedParameters.QueueSize == 123 for request in captured)
    assert all(
        request.RequestedParameters.Filter.Trigger == ua.DataChangeTrigger.StatusValueTimestamp
        for request in captured
    )


def test_abrupt_process_exit_recovers_committed_record_even_without_raw_flush(tmp_path):
    child = r"""
import asyncio, os, sys
from dataclasses import replace
from pathlib import Path
from datetime import UTC, datetime
from edge_logger.config import load_config
from edge_logger.models import ProcessRecord
from edge_logger.runtime import LoggerRuntime
from edge_logger.storage.audit_db import AuditDatabase
p=Path(sys.argv[1])
c=replace(load_config(Path('config/config.example.yaml')),raw_path=p/'raw',reports_path=p/'reports',state_path=p/'state')
d=AuditDatabase(c.state_path/'state.sqlite3')
r=LoggerRuntime(c,d)
t=datetime.now(UTC)
asyncio.run(r.ingest_record(ProcessRecord('DEMO-01',t,t,'1','opcua:cycle'),0))
os._exit(0)
"""
    subprocess.run([sys.executable, "-c", child, str(tmp_path)], check=True)
    config = _config(tmp_path)
    database = AuditDatabase(config.state_path / "state.sqlite3")
    try:
        runtime = LoggerRuntime(config, database)
        asyncio.run(runtime.recover())
        period = database.open_period()
        segments = database.raw_segments_for_period(period["id"])
        assert period["record_count"] == 1 and len(segments) == 1
        assert [row.cycle_id for row in read_records(Path(segments[0]["path"]))] == ["1"]
        asyncio.run(runtime.recover())
        assert database.open_period()["record_count"] == 1
        assert len(database.raw_segments_for_period(period["id"])) == 1
    finally:
        database.close()


def test_canonical_record_survives_raw_write_failure(tmp_path, monkeypatch):
    config = _config(tmp_path)
    database = AuditDatabase(config.state_path / "state.sqlite3")
    runtime = LoggerRuntime(config, database)
    original = RawArchiveWriter.append
    monkeypatch.setattr(
        RawArchiveWriter,
        "append",
        lambda *_: (_ for _ in ()).throw(OSError("disk mirror interrupted")),
    )
    with pytest.raises(OSError):
        asyncio.run(runtime.ingest_record(_record(), 1))
    monkeypatch.setattr(RawArchiveWriter, "append", original)
    asyncio.run(runtime._close_active_segment())
    asyncio.run(runtime.recover())
    period = database.open_period()
    archive = Path(database.raw_segments_for_period(period["id"])[0]["path"])
    assert len(list(read_records(archive))) == period["record_count"] == 1
    database.close()


def test_interrupted_import_resumes_and_overlapping_exports_do_not_duplicate(tmp_path, monkeypatch):
    config = _config(tmp_path)
    source = tmp_path / "source.csv"
    content = (
        "Date\tTime\tCycles\tCycle time\n24.10.2026\t10:00:00\t1\t10\n24.10.2026\t10:00:10\t2\t10\n"
    )
    source.write_text(content, encoding="utf-16")
    database = AuditDatabase(config.state_path / "state.sqlite3")
    runtime = LoggerRuntime(config, database)
    original = FileExportImporter.records

    def interrupted(self, path):
        yield next(original(self, path))
        raise OSError("import interruption")

    monkeypatch.setattr(FileExportImporter, "records", interrupted)
    with pytest.raises(OSError):
        asyncio.run(runtime.import_file(source))
    monkeypatch.setattr(FileExportImporter, "records", original)
    result = asyncio.run(runtime.import_file(source))
    assert not result.duplicate and result.imported_records == 1
    assert database.open_period()["record_count"] == 2
    overlap = tmp_path / "overlap.csv"
    overlap.write_text(content + "\n", encoding="utf-16")
    assert asyncio.run(runtime.import_file(overlap)).imported_records == 0
    assert database.open_period()["record_count"] == 2
    assert asyncio.run(runtime.import_file(source)).duplicate
    database.close()


def test_healthy_clock_excludes_outage_and_reconnect_interval(tmp_path, monkeypatch):
    config = _config(tmp_path)
    database = AuditDatabase(config.state_path / "state.sqlite3")
    runtime = LoggerRuntime(config, database)
    provider = SimpleNamespace(
        connected=True, healthy_at=0.0, connection_epoch=1, production_running=True
    )
    runtime._provider = provider
    clock = [0.0]
    states = iter([(1, True, 1), (2, True, 1), (7202, False, 1), (7203, True, 2), (7204, True, 2)])

    async def advance(_):
        try:
            clock[0], provider.connected, provider.connection_epoch = next(states)
            provider.healthy_at = clock[0]
        except StopIteration:
            runtime.stop_event.set()

    monkeypatch.setattr(runtime_module, "time", SimpleNamespace(monotonic=lambda: clock[0]))
    monkeypatch.setattr(runtime, "_pause", advance)
    asyncio.run(runtime._clock_loop())
    assert database.open_period()["accumulated_seconds"] == 3
    database.close()


def test_live_run_sends_newly_generated_report_and_preserves_acquisition(tmp_path, monkeypatch):
    config = _config(tmp_path)
    email = {
        "enabled": True,
        "host": "test.invalid",
        "sender": "logger@test.invalid",
        "recipients": ["qa@test.invalid"],
    }
    config = replace(
        config,
        reporting=replace(config.reporting, acquisition_hours=0.000001),
        raw={**config.raw, "email": email},
    )
    database = AuditDatabase(config.state_path / "state.sqlite3")
    runtime = LoggerRuntime(config, database)
    sent = []

    class Provider(AcquisitionProvider):
        connected = True
        connection_epoch = 1
        production_running = True

        @property
        def healthy_at(self):
            return time.monotonic()

        async def records(self):
            yield _record()
            await runtime.stop_event.wait()

    async def exercise():
        loop = asyncio.get_running_loop()

        def send(_, report_id, workbook, description, records):
            assert workbook.exists() and records.exists()
            sent.append(report_id)
            loop.call_soon_threadsafe(runtime.stop_event.set)
            return f"<{report_id}@edge-logger.local>"

        async def quick_pause(_):
            await asyncio.sleep(0.01)

        monkeypatch.setattr("edge_logger.email.outbox.send_report", send)
        monkeypatch.setattr(runtime, "_pause", quick_pause)
        await asyncio.wait_for(runtime.run(Provider()), 10)

    asyncio.run(exercise())
    assert len(sent) == 1
    assert database.email_counts_by_state() == {"SENT": 1}
    database.close()


def test_export_receiver_is_idempotent_and_requires_exact_storage_acknowledgement(tmp_path):
    config = _config(tmp_path)
    database = AuditDatabase(config.state_path / "state.sqlite3")
    runtime = LoggerRuntime(config, database)
    asyncio.run(runtime.ingest_record(_record(), 0))
    asyncio.run(runtime._close_active_segment())
    period = database.force_seal_open_period()
    export = runtime.export_period(period["report_id"])
    receiver = tmp_path / "receiver.sqlite3"
    receipt = ingest_export(Path(export["path"]), receiver)
    assert receipt["record_count"] == 1
    assert ingest_export(Path(export["path"]), receiver)["duplicate"]
    workbook = config.reports_path / "report.xlsx"
    workbook.write_bytes(b"report")
    database.set_workbook(period["id"], workbook, archive_checksum(workbook))
    with pytest.raises(ValueError, match="process export"):
        database.acknowledge_storage(
            period["report_id"], archive_checksum(workbook), receipt["receipt"], "wrong"
        )
    database.acknowledge_storage(
        period["report_id"], archive_checksum(workbook), receipt["receipt"], export["sha256"]
    )
    assert database.period(period["id"])["storage_acknowledged_at"]
    with gzip.open(export["path"], "rt") as stream:
        entry = json.loads(stream.readline())
    entry["record"]["cycle_id"] = "changed"
    conflicting = tmp_path / "conflict.gz"
    with gzip.open(conflicting, "wt") as stream:
        stream.write(json.dumps(entry) + "\n")
    with pytest.raises(ValueError, match="Conflicting payload"):
        ingest_export(conflicting, receiver)
    database.close()


def test_partial_recipient_refusal_is_not_reported_as_success(monkeypatch):
    class SMTP:
        def __init__(self, *_args, **_kwargs):
            pass

        def __enter__(self):
            return self

        def __exit__(self, *_):
            pass

        def send_message(self, _):
            return {"failed@test.invalid": (550, b"refused")}

    monkeypatch.setattr("edge_logger.email.sender.smtplib.SMTP", SMTP)
    with pytest.raises(PartialDeliveryError):
        _send_message(
            EmailConfiguration(
                "test.invalid",
                25,
                "sender@test.invalid",
                ("accepted@test.invalid", "failed@test.invalid"),
                tls_mode="none",
            ),
            EmailMessage(),
        )


def test_acquisition_lock_excludes_second_process_and_releases(tmp_path):
    with (
        AcquisitionLock(tmp_path),
        pytest.raises(RuntimeError, match="Another logger"),
        AcquisitionLock(tmp_path),
    ):
        pass
    with AcquisitionLock(tmp_path):
        pass


def test_verified_ups_utility_and_durable_shutdown_request(tmp_path):
    command = [
        sys.executable,
        "-c",
        'print(\'{"external_power_present":false,"battery_critical":true,"battery_percent":3}\')',
    ]
    status = asyncio.run(CommandUpsProvider(command).status())
    assert (
        not status.external_power_present
        and status.battery_critical
        and status.battery_percent == 3
    )
    asyncio.run(request_host_shutdown(tmp_path))
    assert json.loads((tmp_path / "shutdown-request.json").read_text())["reason"]


def test_multicavity_reject_quantity_is_distinct_from_bad_cycle_indicator():
    record = replace(
        _record(),
        production_context={"active_cavities": 4, "sample_parts": 0},
        signals=(
            SignalValue(
                "bad", "Rejects", 2, None, "int", QualityStatus.GOOD, _record().source_timestamp
            ),
        ),
    )
    registry = {"bad": {"exception_rule": "bad_part", "bad_part_mode": "quantity"}}
    analysis = analyse_records(lambda: iter([record]), registry, "Europe/London", 5, {})
    assert analysis.bad_part_count == 1
    assert (
        analysis.produced_parts == 4 and analysis.rejected_parts == 2 and analysis.good_parts == 2
    )
    registry["bad"]["bad_part_mode"] = "indicator"
    analysis = analyse_records(lambda: iter([record]), registry, "Europe/London", 5, {})
    assert analysis.rejected_parts is None and analysis.good_parts is None


@pytest.mark.parametrize("field", ["acquisition_hours", "maximum_calendar_days"])
def test_configuration_rejects_nonfinite_reporting_thresholds(tmp_path, field):
    source = Path("config/config.example.yaml").read_text()
    value = "120" if field == "acquisition_hours" else "null"
    source = source.replace(f"{field}: {value}", f"{field}: .nan")
    path = tmp_path / "config.yaml"
    path.write_text(source)
    with pytest.raises(ConfigurationError):
        load_config(path)


def test_record_cap_seals_and_opens_successor_without_waiting_for_hours(tmp_path):
    config = _config(tmp_path)
    config = replace(config, reporting=replace(config.reporting, maximum_records=2))
    database = AuditDatabase(config.state_path / "state.sqlite3")
    runtime = LoggerRuntime(config, database)
    assert asyncio.run(runtime.ingest_record(_record(1), 0)) is None
    sealed = asyncio.run(runtime.ingest_record(_record(2), 0))
    assert sealed["record_count"] == 2 and sealed["accumulated_seconds"] == 0
    asyncio.run(runtime.ingest_record(_record(3), 0))
    assert database.open_period()["record_count"] == 1
    asyncio.run(runtime._close_active_segment())
    database.close()


def test_calendar_limit_seals_nonempty_period_without_healthy_hours(tmp_path):
    config = _config(tmp_path)
    config = replace(config, reporting=replace(config.reporting, maximum_calendar_days=1))
    database = AuditDatabase(config.state_path / "state.sqlite3")
    runtime = LoggerRuntime(config, database)
    asyncio.run(runtime.ingest_record(_record(), 0))
    period = database.open_period()
    database.connection.execute(
        "UPDATE report_periods SET started_at = ? WHERE id = ?",
        ((datetime.now(UTC) - timedelta(days=2)).isoformat(), period["id"]),
    )
    asyncio.run(runtime._seal_if_due())
    assert database.period(period["id"])["state"] == "SEALED"
    assert database.open_period()["id"] != period["id"]
    database.close()


def test_smtp_work_does_not_block_next_cycle_commit(tmp_path, monkeypatch):
    config = _config(tmp_path)
    database = AuditDatabase(config.state_path / "state.sqlite3")
    runtime = LoggerRuntime(config, database)
    asyncio.run(runtime.ingest_record(_record(), 0))
    asyncio.run(runtime._close_active_segment())
    period = database.force_seal_open_period()
    database.transition_period(period["id"], ReportPeriodState.REPORT_BUILDING)
    database.transition_period(period["id"], ReportPeriodState.REPORT_VALIDATED)
    workbook = tmp_path / "report.xlsx"
    workbook.write_bytes(b"workbook")
    database.set_workbook(period["id"], workbook, archive_checksum(workbook))
    processor = EmailOutboxProcessor(
        {
            "enabled": True,
            "host": "test.invalid",
            "sender": "logger@test.invalid",
            "recipients": ["qa@test.invalid"],
        },
        database,
    )
    processor.queue(period["id"])
    started = threading.Event()
    release = threading.Event()

    def send(_, report_id, *_args):
        started.set()
        assert release.wait(5)
        return f"<{report_id}@edge-logger.local>"

    monkeypatch.setattr("edge_logger.email.outbox.send_report", send)

    async def exercise():
        sending = asyncio.create_task(processor.process_pending_async())
        try:
            while not started.is_set():
                await asyncio.sleep(0.001)
            await asyncio.wait_for(runtime.ingest_record(_record(2), 0), 0.5)
            assert database.open_period()["record_count"] == 1
        finally:
            release.set()
            await sending
        await runtime._close_active_segment()

    asyncio.run(exercise())
    assert database.email_counts_by_state() == {"SENT": 1}
    database.close()


def test_retry_backoff_skips_pending_attempt_until_due(tmp_path, monkeypatch):
    config = _config(tmp_path)
    database = AuditDatabase(config.state_path / "state.sqlite3")
    runtime = LoggerRuntime(config, database)
    asyncio.run(runtime.ingest_record(_record(), 0))
    asyncio.run(runtime._close_active_segment())
    period = database.force_seal_open_period()
    database.transition_period(period["id"], ReportPeriodState.REPORT_BUILDING)
    database.transition_period(period["id"], ReportPeriodState.REPORT_VALIDATED)
    workbook = tmp_path / "report.xlsx"
    workbook.write_bytes(b"workbook")
    database.set_workbook(period["id"], workbook, archive_checksum(workbook))
    processor = EmailOutboxProcessor(
        {
            "enabled": True,
            "host": "test.invalid",
            "sender": "logger@test.invalid",
            "recipients": ["qa@test.invalid"],
        },
        database,
    )
    processor.queue(period["id"])
    monkeypatch.setattr(
        "edge_logger.email.outbox.send_report",
        lambda *_: (_ for _ in ()).throw(ConnectionRefusedError("offline")),
    )
    assert processor.process_pending()[0].state is EmailState.FAILED_RETRYABLE
    assert processor.process_pending() == ()
    database.connection.execute(
        "UPDATE email_outbox SET next_attempt_at = ?",
        ((datetime.now(UTC) - timedelta(seconds=1)).isoformat(),),
    )
    assert len(processor.process_pending()) == 1
    assert database.email_outbox(period["id"])["attempt_count"] == 2
    database.close()


@pytest.mark.parametrize(
    "rows",
    [
        "Date\tTime\tCycles\tPressure\t\n24.10.26\t12:00:00\t1\t100\tUNNAMED\n",
        "Date\tTime\tCycles\tPressure\t Pressure\n24.10.26\t12:00:00\t1\t100\t200\n",
    ],
)
def test_parser_refuses_ambiguous_or_unnamed_data_columns(tmp_path, rows):
    source = tmp_path / "malformed.csv"
    source.write_text(rows, encoding="utf-16")
    with pytest.raises(ValueError):
        FileExportImporter("DEMO-01", {}).inspect(source)


def test_parser_flags_nonfinite_values_and_ambiguous_dst_times(tmp_path):
    source = tmp_path / "source.csv"
    source.write_text(
        "Date\tTime\tCycles\tPressure\n25.10.26\t01:30:00\t1\tNaN\n", encoding="utf-16"
    )
    result = next(
        FileExportImporter(
            "DEMO-01", {"p": {"source_name": "Pressure", "expected_datatype": "float"}}
        ).records(source)
    )
    assert result.source_timestamp is None
    assert result.signals[0].quality is QualityStatus.INVALID_NUMERIC


def test_encoded_email_size_is_checked_before_smtp(monkeypatch):
    message = EmailMessage()
    message.set_content("large body" * 100)
    configuration = EmailConfiguration(
        "test.invalid", 25, "logger@test.invalid", ("qa@test.invalid",), max_message_bytes=10
    )
    monkeypatch.setattr(
        "edge_logger.email.sender.smtplib.SMTP",
        lambda *_args, **_kwargs: pytest.fail("network must not be used"),
    )
    from edge_logger.email.sender import AttachmentTooLargeError

    with pytest.raises(AttachmentTooLargeError):
        _send_message(configuration, message)


def test_report_preserves_job_context_as_literal_text_and_respects_run_resets(tmp_path):
    from openpyxl import load_workbook

    from edge_logger.reporting.generator import generate_report

    records = [
        replace(
            _record(1), run_id="boot-1", production_context={"job_id": "=1+1", "active_cavities": 4}
        ),
        replace(
            _record(2),
            cycle_id="1",
            run_id="boot-2",
            production_context={"job_id": "=1+1", "active_cavities": 4},
        ),
    ]
    analysis = analyse_records(lambda: iter(records), {}, "Europe/London", 5, {})
    assert (
        analysis.unique_cycles == 2
        and analysis.duplicate_cycles == 0
        and analysis.missing_cycles == 0
    )
    path, _ = generate_report(records, tmp_path, "context-test", "DEMO-01", "Europe/London", {})
    book = load_workbook(path, read_only=True)
    try:
        raw = book["Active Raw Data"]
        headers = [cell.value for cell in next(raw.iter_rows())]
        job_column = headers.index("Job ID") + 1
        assert raw.cell(2, job_column).value == "=1+1"
        assert raw.cell(2, job_column).data_type == "s"
        assert raw.cell(2, headers.index("Run ID") + 1).value == "boot-1"
    finally:
        book.close()


def test_existing_v5_database_is_backed_up_and_migrated_without_recounting_cycles(tmp_path):
    from edge_logger.storage.source_archive import archive_source_file

    config = _config(tmp_path)
    source = tmp_path / "source.csv"
    source.write_text(
        "Date\tTime\tCycles\tPressure\n24.10.26\t13:00:01\t1\t101\n", encoding="utf-16"
    )
    archived = archive_source_file(source, config.raw_path / "source")
    path = config.state_path / "state.sqlite3"
    database = AuditDatabase(path)
    period = database.get_or_create_open_period()
    database.increment_record_count()
    with RawArchiveWriter(config.raw_path / "records", f"{period['id']}-legacy") as writer:
        writer.append(_record())
    database.register_raw_segment(period["id"], writer.path, writer.checksum, 1)
    database.register_source_file(
        archived.sha256, source.name, source.stat().st_size, archived.archive_path, period["id"]
    )
    database.associate_source_with_period(archived.sha256, period["id"])
    database.close()
    # Restore precisely the pre-change schema while retaining real source/raw evidence.
    with sqlite3.connect(path) as legacy:
        legacy.execute("DROP TABLE process_records")
        legacy.execute("DROP INDEX one_open_period")
        for table, columns in {
            "source_files": ("import_state", "imported_records"),
            "report_periods": ("storage_acknowledged_at", "storage_receipt"),
            "email_outbox": ("next_attempt_at",),
        }.items():
            for column in columns:
                legacy.execute(f"ALTER TABLE {table} DROP COLUMN {column}")
        legacy.execute("DELETE FROM schema_migrations WHERE version >= 6")
    database = AuditDatabase(path)
    try:
        backup = path.with_name(f"{path.name}.pre-v6-backup")
        assert backup.is_file() and AuditDatabase.read_only_check(backup)["valid"]
        assert not database.import_complete(archived.sha256)
        runtime = LoggerRuntime(config, database)
        asyncio.run(runtime.recover())
        assert (
            database.connection.execute("SELECT COUNT(*) FROM process_records").fetchone()[0] == 1
        )
        assert database.open_period()["record_count"] == 1
        result = asyncio.run(runtime.import_file(source))
        assert not result.duplicate and result.imported_records == 0
        assert database.import_complete(archived.sha256)
        assert database.open_period()["record_count"] == 1
    finally:
        database.close()
